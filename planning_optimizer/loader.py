from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
import json
from pathlib import Path
from typing import Any, BinaryIO, Iterable
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from openpyxl import load_workbook
from jsonschema import Draft202012Validator, FormatChecker

from .calendars import build_calendar_slots, contiguous_runs, parse_datetime


TABLE_SHEETS = {
    "Systems": ("systems", "system_id"),
    "Packages": ("packages", "package_id"),
    "Activities": ("activities", "activity_id"),
    "Gates": ("gates", "gate_id"),
    "Dependencies": ("dependencies", "dependency_id"),
    "Resources": ("resources", "resource_id"),
    "ActivityResources": ("activity_resources", "activity_id"),
    "ResourceSubstitutions": ("resource_substitutions", "required_resource_id"),
    "Zones": ("zones", "zone_id"),
    "ActivityZones": ("activity_zones", "activity_id"),
    "Calendars": ("calendars", "calendar_id"),
    "CalendarShifts": ("calendar_shifts", "calendar_id"),
    "MilestonePriorities": ("milestone_priorities", "gate_id"),
}
OPTIONAL_SHEETS = {"ResourceSubstitutions"}
SCHEMA_PATH = Path(__file__).resolve().parents[1] / "schema" / "planning_optimizer_schema_v1.json"

BOOL_FIELDS = {
    "enabled",
    "preemptible",
    "exposed",
    "is_project_milestone",
    "unlimited",
    "exclusive",
    "is_override",
    "requires_system_arrival",
}


@dataclass
class ValidationReport:
    errors: list[str]
    warnings: list[str]

    @property
    def ok(self) -> bool:
        return not self.errors

    def as_dict(self) -> dict[str, Any]:
        return {"ok": self.ok, "errors": self.errors, "warnings": self.warnings}


@dataclass
class ProjectData:
    data: dict[str, Any]
    report: ValidationReport


class ProjectValidationError(ValueError):
    def __init__(self, report: ValidationReport):
        super().__init__("Project validation failed")
        self.report = report


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _coerce_bool(value: Any, location: str, warnings: list[str]) -> Any:
    if isinstance(value, bool) or value is None:
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        warnings.append(f"{location}: coerced numeric {value!r} to boolean.")
        return bool(value)
    if isinstance(value, str) and value.strip().lower() in {"true", "false"}:
        warnings.append(f"{location}: coerced text {value!r} to boolean.")
        return value.strip().lower() == "true"
    return value


def _json_scalar(value: Any) -> Any:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return value.isoformat()
    return value


def _find_header_row(ws, key_header: str) -> tuple[int, list[str]]:
    for row_number, row in enumerate(ws.iter_rows(min_row=1, max_row=30, values_only=True), start=1):
        values = [str(v).strip() if v is not None else "" for v in row]
        if key_header in values:
            return row_number, values
    raise ValueError(f"Sheet {ws.title!r}: could not find header {key_header!r}.")


def _read_table(ws, key_header: str, warnings: list[str]) -> list[dict[str, Any]]:
    header_row, headers = _find_header_row(ws, key_header)
    last_column = max(i for i, header in enumerate(headers) if header)
    headers = headers[: last_column + 1]
    rows: list[dict[str, Any]] = []
    for excel_row, values in enumerate(
        ws.iter_rows(min_row=header_row + 1, max_col=len(headers), values_only=True),
        start=header_row + 1,
    ):
        if all(_is_blank(v) for v in values):
            continue
        record: dict[str, Any] = {}
        for header, value in zip(headers, values):
            if not header:
                continue
            if header in BOOL_FIELDS:
                value = _coerce_bool(value, f"{ws.title}!{header} row {excel_row}", warnings)
            # Empty optional cells are absent in the normalized model rather
            # than being converted to JSON null indiscriminately.
            if not _is_blank(value):
                record[header] = _json_scalar(value)
        rows.append(record)
    return rows


def _read_metadata(ws, warnings: list[str]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for row_number, row in enumerate(ws.iter_rows(min_row=2, max_col=2, values_only=True), start=2):
        field, value = row
        if _is_blank(field):
            continue
        result[str(field).strip()] = _json_scalar(value)
    return result


def _timezone_name(data: dict[str, Any]) -> str:
    active = data.get("metadata", {}).get("active_calendar")
    calendars = {row.get("calendar_id"): row for row in data.get("calendars", [])}
    timezone = calendars.get(active, {}).get("timezone")
    return timezone or "UTC"


def _normalize_datetime(value: Any, timezone_name: str, location: str, errors: list[str]) -> Any:
    if value is None:
        return value
    try:
        zone = ZoneInfo(timezone_name)
    except ZoneInfoNotFoundError:
        errors.append(f"{location}: unknown timezone {timezone_name!r}.")
        zone = ZoneInfo("UTC")
    try:
        if isinstance(value, datetime):
            parsed = value
        else:
            parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=zone)
        return parsed.isoformat()
    except (TypeError, ValueError):
        errors.append(f"{location}: invalid date-time {value!r}.")
        return value


def _normalize_data(data: dict[str, Any], warnings: list[str], errors: list[str]) -> None:
    timezone_name = _timezone_name(data)
    metadata = data.get("metadata", {})
    metadata["project_start"] = _normalize_datetime(
        metadata.get("project_start"), timezone_name, "metadata.project_start", errors
    )
    if isinstance(metadata.get("created_date"), (date, datetime)):
        metadata["created_date"] = metadata["created_date"].date().isoformat()

    for row in data.get("systems", []):
        row["arrival_date"] = _normalize_datetime(
            row.get("arrival_date"), timezone_name, f"system {row.get('system_id')}.arrival_date", errors
        )
    for row in data.get("calendars", []):
        for field in ("valid_from", "valid_to"):
            value = row.get(field)
            if isinstance(value, datetime):
                row[field] = value.date().isoformat()
            elif isinstance(value, date):
                row[field] = value.isoformat()
        weekend = row.get("weekend_days")
        if isinstance(weekend, str):
            row["weekend_days"] = [item.strip().upper() for item in weekend.split(",") if item.strip()]

    for activity in data.get("activities", []):
        activity.setdefault("requires_system_arrival", True)


def _duplicates(records: Iterable[dict[str, Any]], field: str) -> list[Any]:
    seen: set[Any] = set()
    duplicates: set[Any] = set()
    for row in records:
        value = row.get(field)
        if value in seen:
            duplicates.add(value)
        seen.add(value)
    return sorted(duplicates, key=str)


def _check_required(row: dict[str, Any], required: Iterable[str], label: str, errors: list[str]) -> None:
    for field in required:
        if field not in row or row[field] is None or row[field] == "":
            errors.append(f"{label}: missing required field {field!r}.")


def validate_project(data: dict[str, Any], warnings: list[str] | None = None) -> ValidationReport:
    warnings = list(warnings or [])
    errors: list[str] = []
    required_collections = [
        "metadata", "systems", "packages", "activities", "gates", "dependencies",
        "resources", "activity_resources", "zones", "activity_zones", "calendars",
        "calendar_shifts", "milestone_priorities",
    ]
    for key in required_collections:
        if key not in data:
            errors.append(f"Missing top-level collection {key!r}.")

    # Structural constraints share the same V1 schema as the browser validator.
    schema = json.loads(SCHEMA_PATH.read_text(encoding="utf-8"))
    schema_errors = Draft202012Validator(schema, format_checker=FormatChecker()).iter_errors(data)
    errors.extend(f"JSON Schema {error.json_path}: {error.message}" for error in schema_errors)

    metadata = data.get("metadata", {})
    _check_required(metadata, ["project_id", "revision_id", "project_start"], "metadata", errors)

    id_fields = {
        "systems": "system_id", "packages": "package_id", "activities": "activity_id",
        "gates": "gate_id", "dependencies": "dependency_id", "resources": "resource_id",
        "zones": "zone_id", "calendars": "calendar_id",
    }
    for collection, field in id_fields.items():
        duplicates = _duplicates(data.get(collection, []), field)
        if duplicates:
            errors.append(f"{collection}: duplicate {field} values {duplicates}.")

    systems = {r.get("system_id"): r for r in data.get("systems", [])}
    packages = {r.get("package_id"): r for r in data.get("packages", [])}
    activities = {r.get("activity_id"): r for r in data.get("activities", [])}
    gates = {r.get("gate_id"): r for r in data.get("gates", [])}
    resources = {r.get("resource_id"): r for r in data.get("resources", [])}
    zones = {r.get("zone_id"): r for r in data.get("zones", [])}
    calendars = {r.get("calendar_id"): r for r in data.get("calendars", [])}
    if metadata.get("objective_gate") and metadata["objective_gate"] not in gates:
        errors.append(f"metadata: unknown objective_gate {metadata['objective_gate']!r}.")
    for system_id, row in systems.items():
        _check_required(row, ["system_id", "name", "family", "arrival_date", "enabled"], f"system {system_id}", errors)
        if row.get("installation_zone") and row["installation_zone"] not in zones:
            errors.append(f"system {system_id}: unknown installation_zone {row['installation_zone']!r}.")
        if not isinstance(row.get("enabled"), bool):
            errors.append(f"system {system_id}: enabled must be boolean.")

    for package_id, row in packages.items():
        _check_required(row, ["package_id", "system_id", "name", "enabled"], f"package {package_id}", errors)
        if row.get("system_id") not in systems:
            errors.append(f"package {package_id}: unknown system {row.get('system_id')!r}.")

    for activity_id, row in activities.items():
        _check_required(
            row,
            ["activity_id", "package_id", "name", "duration_h", "duration_basis", "preemptible", "enabled"],
            f"activity {activity_id}",
            errors,
        )
        if row.get("package_id") not in packages:
            errors.append(f"activity {activity_id}: unknown package {row.get('package_id')!r}.")
        if row.get("duration_basis") not in {"WORK_TIME", "ELAPSED_TIME"}:
            errors.append(f"activity {activity_id}: invalid duration_basis {row.get('duration_basis')!r}.")
        duration = row.get("duration_h")
        if isinstance(duration, bool) or not isinstance(duration, (int, float)) or duration <= 0 or int(duration) != duration:
            errors.append(f"activity {activity_id}: duration_h must be a positive whole number of hours.")
        if not isinstance(row.get("preemptible"), bool):
            errors.append(f"activity {activity_id}: preemptible must be boolean.")
        if not isinstance(row.get("enabled"), bool):
            errors.append(f"activity {activity_id}: enabled must be boolean.")
        if not isinstance(row.get("requires_system_arrival", True), bool):
            errors.append(f"activity {activity_id}: requires_system_arrival must be boolean.")
        if row.get("duration_basis") == "ELAPSED_TIME" and row.get("preemptible"):
            errors.append(f"activity {activity_id}: ELAPSED_TIME cannot be preemptible in V1.")
        calendar_id = row.get("calendar_id") or metadata.get("active_calendar")
        if row.get("duration_basis") == "WORK_TIME" and calendar_id not in calendars:
            errors.append(f"activity {activity_id}: unknown effective calendar {calendar_id!r}.")
        if row.get("duration_basis") == "WORK_TIME" and row.get("preemptible") is False and calendar_id in calendars:
            relevant_calendars = [calendar_id]
            for demand in data.get("activity_resources", []):
                if demand.get("activity_id") == activity_id and demand.get("resource_id") in resources:
                    required_calendar = resources[demand["resource_id"]].get("calendar_id")
                    if required_calendar:
                        relevant_calendars.append(required_calendar)
            for demand in data.get("activity_zones", []):
                if demand.get("activity_id") == activity_id and demand.get("zone_id") in zones:
                    required_calendar = zones[demand["zone_id"]].get("calendar_id")
                    if required_calendar:
                        relevant_calendars.append(required_calendar)
            try:
                sample_horizon_h = 14 * 24
                slots = build_calendar_slots(data, parse_datetime(metadata["project_start"]), sample_horizon_h)
                effective = [
                    all(slots[required_calendar][hour] for required_calendar in set(relevant_calendars))
                    for hour in range(sample_horizon_h)
                ]
                max_continuous_h = max((end - start for start, end in contiguous_runs(effective)), default=0)
                if float(row.get("duration_h", 0)) > max_continuous_h + 1e-9:
                    errors.append(
                        f"activity {activity_id}: non-interruptible duration {row.get('duration_h')} h exceeds "
                        f"the {max_continuous_h:g} h maximum continuous window of its effective calendars."
                    )
            except (KeyError, TypeError, ValueError):
                # Detailed calendar/reference errors are reported separately.
                pass

    for gate_id, row in gates.items():
        _check_required(row, ["gate_id", "name", "gate_type", "exposed"], f"gate {gate_id}", errors)
        package_id = row.get("package_id")
        system_id = row.get("system_id")
        if package_id and package_id not in packages:
            errors.append(f"gate {gate_id}: unknown package {package_id!r}.")
        if system_id and system_id not in systems:
            errors.append(f"gate {gate_id}: unknown system {system_id!r}.")
        if package_id and system_id and packages.get(package_id, {}).get("system_id") != system_id:
            errors.append(f"gate {gate_id}: package/system ownership is inconsistent.")

    completion_gates = [gate_id for gate_id, row in gates.items() if row.get("gate_type") == "PROJECT_COMPLETE"]
    if len(completion_gates) != 1:
        errors.append(f"Exactly one PROJECT_COMPLETE gate is required; found {len(completion_gates)}.")

    for row in data.get("resources", []):
        resource_id = row.get("resource_id")
        _check_required(row, ["resource_id", "name", "type", "capacity", "unlimited"], f"resource {resource_id}", errors)
        if row.get("type") not in {"HUMAN", "EQUIPMENT", "WORKFRONT"}:
            errors.append(f"resource {resource_id}: invalid type {row.get('type')!r}.")
        if not isinstance(row.get("capacity"), (int, float)) or row.get("capacity", -1) < 0:
            errors.append(f"resource {resource_id}: capacity must be non-negative.")
        if row.get("calendar_id") and row["calendar_id"] not in calendars:
            errors.append(f"resource {resource_id}: unknown calendar {row['calendar_id']!r}.")

    for row in data.get("zones", []):
        zone_id = row.get("zone_id")
        _check_required(row, ["zone_id", "name", "capacity"], f"zone {zone_id}", errors)
        if not isinstance(row.get("capacity"), (int, float)) or row.get("capacity", 0) <= 0:
            errors.append(f"zone {zone_id}: capacity must be positive.")
        if row.get("calendar_id") and row["calendar_id"] not in calendars:
            errors.append(f"zone {zone_id}: unknown calendar {row['calendar_id']!r}.")

    seen_pairs: set[tuple[Any, Any]] = set()
    for row in data.get("activity_resources", []):
        pair = (row.get("activity_id"), row.get("resource_id"))
        if pair in seen_pairs:
            errors.append(f"activity_resources: duplicate mapping {pair}.")
        seen_pairs.add(pair)
        if pair[0] not in activities:
            errors.append(f"activity_resources: unknown activity {pair[0]!r}.")
        if pair[1] not in resources:
            errors.append(f"activity_resources: unknown resource {pair[1]!r}.")
        quantity = row.get("quantity")
        if not isinstance(quantity, (int, float)) or quantity <= 0:
            errors.append(f"activity_resources {pair}: quantity must be positive.")
        elif pair[1] in resources and not resources[pair[1]].get("unlimited") and quantity > resources[pair[1]].get("capacity", 0):
            errors.append(f"activity_resources {pair}: demand {quantity} exceeds capacity {resources[pair[1]].get('capacity')}.")

    for row in data.get("resource_substitutions", []):
        for field in ("required_resource_id", "substitute_resource_id"):
            if row.get(field) not in resources:
                errors.append(f"resource_substitutions: unknown {field} {row.get(field)!r}.")

    seen_pairs.clear()
    for row in data.get("activity_zones", []):
        pair = (row.get("activity_id"), row.get("zone_id"))
        if pair in seen_pairs:
            errors.append(f"activity_zones: duplicate mapping {pair}.")
        seen_pairs.add(pair)
        if pair[0] not in activities:
            errors.append(f"activity_zones: unknown activity {pair[0]!r}.")
        if pair[1] not in zones:
            errors.append(f"activity_zones: unknown zone {pair[1]!r}.")
        if row.get("load") == "ALL" and not row.get("exclusive"):
            errors.append(f"activity_zones {pair}: load ALL requires exclusive=true.")
        if row.get("exclusive") and row.get("load") != "ALL":
            warnings.append(f"activity_zones {pair}: exclusive=true overrides numeric load.")

    nodes = {("ACTIVITY", key) for key in activities} | {("GATE", key) for key in gates}
    enabled_edges: list[tuple[tuple[str, str], tuple[str, str], str]] = []
    for row in data.get("dependencies", []):
        dep_id = row.get("dependency_id")
        source = (row.get("source_type"), row.get("source_id"))
        target = (row.get("target_type"), row.get("target_id"))
        if source not in nodes:
            errors.append(f"dependency {dep_id}: unknown source {source}.")
        if target not in nodes:
            errors.append(f"dependency {dep_id}: unknown target {target}.")
        if row.get("relation") != "FS":
            errors.append(f"dependency {dep_id}: only FS is supported in V1.")
        if not isinstance(row.get("lag_h"), (int, float)) or row.get("lag_h", -1) < 0:
            errors.append(f"dependency {dep_id}: lag_h must be non-negative.")
        if row.get("enabled") is True and source in nodes and target in nodes:
            enabled_edges.append((source, target, dep_id))

    indegree = {node: 0 for node in nodes}
    adjacency: dict[tuple[str, str], list[tuple[str, str]]] = {node: [] for node in nodes}
    for source, target, _ in enabled_edges:
        adjacency[source].append(target)
        indegree[target] += 1
    queue = [node for node, degree in indegree.items() if degree == 0]
    visited = 0
    while queue:
        node = queue.pop()
        visited += 1
        for target in adjacency[node]:
            indegree[target] -= 1
            if indegree[target] == 0:
                queue.append(target)
    if visited != len(nodes):
        cyclic = sorted(f"{kind}:{identifier}" for (kind, identifier), degree in indegree.items() if degree > 0)
        errors.append(f"Enabled dependency graph contains a cycle involving {cyclic}.")

    shift_keys: set[tuple[Any, Any, Any]] = set()
    for row in data.get("calendar_shifts", []):
        key = (row.get("calendar_id"), row.get("weekday"), row.get("shift_name"))
        if key in shift_keys:
            errors.append(f"calendar_shifts: duplicate shift {key}.")
        shift_keys.add(key)
        if row.get("calendar_id") not in calendars:
            errors.append(f"calendar_shifts: unknown calendar {row.get('calendar_id')!r}.")
        for field in ("start_time", "end_time"):
            value = row.get(field)
            try:
                datetime.strptime(str(value), "%H:%M")
            except ValueError:
                errors.append(f"calendar_shifts {key}: invalid {field} {value!r}.")

    seen_priority_gates: set[Any] = set()
    seen_priorities: set[Any] = set()
    enabled_priority_gates: set[Any] = set()
    for row in data.get("milestone_priorities", []):
        gate_id = row.get("gate_id")
        priority = row.get("priority")
        if gate_id in seen_priority_gates:
            errors.append(f"milestone_priorities: duplicate gate {gate_id!r}.")
        seen_priority_gates.add(gate_id)
        if gate_id not in gates:
            errors.append(f"milestone_priorities: unknown gate {gate_id!r}.")
        if row.get("enabled") and (not isinstance(priority, int) or isinstance(priority, bool) or priority < 1):
            errors.append(f"milestone_priorities {gate_id!r}: enabled priority must be a positive integer.")
        elif row.get("enabled"):
            if priority in seen_priorities:
                errors.append(f"milestone_priorities: enabled priority {priority} is duplicated.")
            seen_priorities.add(priority)
            enabled_priority_gates.add(gate_id)
    if len(completion_gates) == 1 and completion_gates[0] not in enabled_priority_gates:
        errors.append("The PROJECT_COMPLETE gate must have an enabled milestone priority.")

    return ValidationReport(errors=errors, warnings=warnings)


def load_project(source: str | Path | BinaryIO, strict: bool = True) -> ProjectData:
    warnings: list[str] = []
    errors: list[str] = []
    if hasattr(source, "seek"):
        source.seek(0)
        workbook_source = source
    else:
        workbook_source = Path(source)
    workbook = load_workbook(workbook_source, read_only=True, data_only=False)
    missing_sheets = [name for name in ["Metadata", *TABLE_SHEETS] if name not in workbook.sheetnames and name not in OPTIONAL_SHEETS]
    if missing_sheets:
        report = ValidationReport(errors=[f"Missing worksheets: {missing_sheets}."], warnings=[])
        if strict:
            raise ProjectValidationError(report)
        return ProjectData(data={}, report=report)

    data: dict[str, Any] = {"metadata": _read_metadata(workbook["Metadata"], warnings)}
    for sheet_name, (key, key_header) in TABLE_SHEETS.items():
        data[key] = _read_table(workbook[sheet_name], key_header, warnings) if sheet_name in workbook.sheetnames else []
    _normalize_data(data, warnings, errors)
    report = validate_project(data, warnings)
    report.errors[:0] = errors
    if strict and not report.ok:
        raise ProjectValidationError(report)
    return ProjectData(data=data, report=report)


def save_normalized_json(project: ProjectData, path: str | Path) -> None:
    Path(path).write_text(json.dumps(project.data, indent=2, ensure_ascii=False), encoding="utf-8")
